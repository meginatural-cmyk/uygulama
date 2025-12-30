import tkinter as tk
from tkinter import filedialog, ttk
import webview
from PIL import Image, ImageTk
import docx
import openpyxl
import os

class AkilliGörüntüleyici:
    def __init__(self, root):
        self.root = root
        self.root.title("Akıllı Dosya Görüntüleyici")
        
        # --- PENCERE BOYUTU (Görseldeki gibi standart pencere) ---
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        # Ekranın büyük bir kısmını kaplar ama tam ekran (fullscreen) değildir
        self.root.geometry(f"{int(screen_width*0.85)}x{int(screen_height*0.85)}+100+50")
        
        self.root.configure(bg="#f0f0f0")

        # --- ÜST PANEL (Koyu Lacivert) ---
        self.top_frame = tk.Frame(self.root, bg="#2c3e50", height=120)
        self.top_frame.pack(side="top", fill="x")

        # DOSYA SEÇ VE AÇ BUTONU
        self.btn_open = tk.Button(
            self.top_frame, 
            text="📁 DOSYA SEÇ VE AÇ", 
            command=self.open_file,
            font=("Arial", 16, "bold"),
            bg="#3498db",
            fg="white",
            activebackground="#2980b9",
            activeforeground="white",
            padx=30,
            pady=12,
            relief="raised",
            cursor="hand2"
        )
        self.btn_open.place(relx=0.5, rely=0.5, anchor="center")

        # --- ANA İÇERİK ALANI (Beyaz) ---
        self.main_container = tk.Frame(self.root, bg="white")
        self.main_container.pack(expand=True, fill="both", padx=15, pady=15)

        self.display_label = tk.Label(self.main_container, bg="white")
        self.tree = ttk.Treeview(self.main_container)

    def clear_screen(self):
        """Ekranı yeni içerik için temizler."""
        self.display_label.pack_forget()
        self.tree.pack_forget()
        self.display_label.config(image="", text="")

    def open_file(self):
        file_path = filedialog.askopenfilename(
            title="Dosya Seçin",
            filetypes=[("Tüm Desteklenenler", "*.pdf *.docx *.xlsx *.png *.jpg *.jpeg *.mp3 *.wav *.mp4 *.txt")]
        )
        
        if not file_path:
            return

        self.clear_screen()
        ext = os.path.splitext(file_path)[1].lower()

        # --- CHROMIUM TABANLI GÖRÜNTÜLEME (PDF, Video ve SES) ---
        # Artık ses dosyaları (.mp3, .wav) doğrudan tarayıcıda açılacak
        if ext in ['.pdf', '.mp4', '.webm', '.mp3', '.wav', '.ogg']:
            webview.create_window(os.path.basename(file_path), file_path)
            webview.start()

        # --- EXCEL (Tablo) ---
        elif ext == '.xlsx':
            self.show_excel(file_path)

        # --- FOTOĞRAF ---
        elif ext in ['.jpg', '.png', '.jpeg', '.gif']:
            self.show_image(file_path)

        # --- WORD / METİN ---
        elif ext in ['.docx', '.txt']:
            self.show_text(file_path, ext)

    def show_excel(self, path):
        self.tree.pack(expand=True, fill="both")
        self.tree.delete(*self.tree.get_children())
        wb = openpyxl.load_workbook(path, data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if rows:
            self.tree["columns"] = [f"c{i}" for i in range(len(rows[0]))]
            self.tree["show"] = "headings"
            for i, col in enumerate(rows[0]):
                self.tree.heading(f"c{i}", text=str(col))
                self.tree.column(f"c{i}", width=120)
            for row in rows[1:]:
                self.tree.insert("", "end", values=row)

    def show_image(self, path):
        self.display_label.pack(expand=True, fill="both")
        img = Image.open(path)
        # Pencere boyutuna göre resmi ölçekle
        self.root.update() # Boyutları doğru almak için güncelle
        img.thumbnail((self.main_container.winfo_width() - 20, self.main_container.winfo_height() - 20))
        self.photo = ImageTk.PhotoImage(img)
        self.display_label.config(image=self.photo, text="")

    def show_text(self, path, ext):
        self.display_label.pack(expand=True, fill="both")
        if ext == '.docx':
            doc = docx.Document(path)
            content = "\n".join([p.text for p in doc.paragraphs])
        else:
            with open(path, "r", encoding="utf-8") as f:
                content = f.read()
        self.display_label.config(text=content, font=("Segoe UI", 11), justify="left", anchor="nw")

if __name__ == "__main__":
    root = tk.Tk()
    app = AkilliGörüntüleyici(root)
    root.mainloop()